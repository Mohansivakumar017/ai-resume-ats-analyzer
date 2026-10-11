import logging
from pathlib import Path
from typing import Any, Dict, Tuple
from pypdf import PdfReader
import docx

from src.services.gemini_service import gemini_service

logger = logging.getLogger(__name__)


def extract_pdf_text(pdf_path: Path) -> str:
    """Extracts raw text from a PDF file."""
    reader = PdfReader(str(pdf_path))
    extracted = []
    for i, page in enumerate(reader.pages):
        page_text = page.extract_text()
        if page_text:
            extracted.append(page_text.strip())
    return "\n\n".join(extracted).strip()


def extract_docx_text(docx_path: Path) -> str:
    """Extracts text from paragraphs and tables in a Word DOCX document."""
    doc = docx.Document(str(docx_path))
    extracted = []

    # Extract standard paragraphs
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            extracted.append(txt)

    # Extract tables (common in resume formatting)
    for table in doc.tables:
        for row in table.rows:
            cell_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cell_texts:
                extracted.append(" | ".join(cell_texts))

    return "\n\n".join(extracted).strip()


def extract_plain_text(txt_path: Path) -> str:
    """Extracts plain text from text-based resume files."""
    with open(txt_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read().strip()


def extract_resume_text(file_path: str) -> str:
    """
    Extracts raw text from PDF, DOCX, or TXT resume files with format detection.
    """
    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"Resume file not found at: {file_path}")

    suffix = path.suffix.lower()

    if suffix == ".pdf":
        text = extract_pdf_text(path)
    elif suffix in (".docx", ".doc"):
        text = extract_docx_text(path)
    elif suffix in (".txt", ".md", ".rtf"):
        text = extract_plain_text(path)
    else:
        # Attempt fallback to docx then plain text
        try:
            text = extract_docx_text(path)
        except Exception:
            text = extract_plain_text(path)

    text = (text or "").strip()
    if not text:
        raise ValueError(
            f"Could not extract readable text from '{path.name}'. "
            "Please ensure the file is not empty or a scanned image."
        )
    return text


def parse_resume(file_path: str) -> Tuple[str, Dict[str, Any]]:
    """
    Parses a candidate's resume from PDF, DOCX, or TXT and returns:
    (raw_text, structured_data_dict)
    """
    raw_text = extract_resume_text(file_path)

    prompt = f"""
You are an expert ATS (Applicant Tracking System) parser.
Extract structured information from the following resume text.

Return ONLY a valid JSON object matching this schema:
{{
    "candidate_name": "string (Full name of candidate found in header/contact)",
    "contact": {{
        "email": "string",
        "phone": "string",
        "location": "string"
    }},
    "skills": ["string"],
    "experience": [
        {{
            "title": "string",
            "organization": "string",
            "years": 0.0,
            "technologies": ["string"],
            "highlights": ["string"]
        }}
    ],
    "education": [
        {{
            "degree": "string",
            "field": "string",
            "institution": "string"
        }}
    ],
    "projects": [
        {{
            "name": "string",
            "technologies": ["string"],
            "description": "string"
        }}
    ],
    "certifications": ["string"]
}}

Resume Text:
{raw_text}
"""

    structured_data = gemini_service.generate_json(
        prompt=prompt,
        system_instruction="You are an unbiased, highly accurate ATS resume parser. Never fabricate missing details."
    )

    return raw_text, structured_data
